#!/usr/bin/env python
"""Reference analysis for the session-4 "para curiosos" Volve exercise.

Implements the rules of the Volve prompt printed on the session-4 page, so the
instructor can check whatever the chatbot returns. Same contract discipline as
dca_referencia.py: if the prompt changes, this script changes with it.

Data: scripts/_cache/volve_production.xlsx (gitignored, never committed — the
Equinor Open Data Licence covers research and study, not redistribution). It is
the official "Volve production data.xlsx" of the 2018 Equinor release, sheet
"Daily Production Data": one row per wellbore-day with ON_STREAM_HRS, average
downhole/wellhead pressure and daily oil/gas/water volumes. Download it from
the official Volve page (https://www.equinor.com/energy/volve-data-sharing) or
a public mirror, and drop it in scripts/_cache/ under that name.

The exercise reads the two producers with the best downhole-pressure coverage:
15/9-F-14 (2008-2016, the full field life) and 15/9-F-11 (2013-2016).

What the prompt asks, and this script mirrors:
  1. Producing day: ON_STREAM_HRS > 0 and BORE_OIL_VOL > 0. Rates normalized
     to 24 h on line. Shut-in day: ON_STREAM_HRS = 0.
  2. Reservoir-pressure proxy: AVG_DOWNHOLE_PRESSURE on shut-in days (> 50 bar
     to drop dead-sensor zeros), monthly median, interpolated to producing
     days. The claim this supports: with water injection the reservoir did NOT
     depressurize while the oil rate collapsed.
  3. Monthly series per well: oil rate, liquid rate, water cut, flowing BHP,
     reservoir-pressure proxy.
  4. Productivity index PI = q_liquid / (pR - pwf) on producing days with both
     pressures; monthly median.
  5. Arps fit on the monthly oil rate from its peak (same log-space fitting as
     dca_referencia.py), for the comparison the page asks: what Arps gets
     right (the trend) and what it cannot see (the mechanism, F-11's cliff).
  6. Decline diagnosis per well: depletion (pR trend), productivity loss (PI
     trend), water (water-cut trend), drawdown management (pR - pwf trend).

    python scripts/volve_referencia.py

Outputs (gitignored, instructor-side):
  scripts/_cache/volve_referencia.xlsx  - monthly series + diagnosis per well
  scripts/_cache/volve_referencia.png   - 4-panel diagnostic per well
"""

from __future__ import annotations

import os
from collections import defaultdict

import numpy as np
from scipy.optimize import curve_fit

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_IN = os.path.join(HERE, '_cache', 'volve_production.xlsx')
OUT_XLSX = os.path.join(HERE, '_cache', 'volve_referencia.xlsx')
OUT_PNG = os.path.join(HERE, '_cache', 'volve_referencia.png')

WELLS = ['15/9-F-14', '15/9-F-11']
P_SENSOR_MIN = 50.0     # bar; below this a "pressure" is a dead sensor
MIN_SHUT_DAYS = 2       # monthly pR proxy needs at least this many shut-in days


def q_arps(t, qi, di, b):
    if b < 1e-6:
        return qi * np.exp(-di * t)
    return qi * (1.0 + b * di * t) ** (-1.0 / b)


def load() -> dict[str, list[dict]]:
    import openpyxl
    if not os.path.exists(XLSX_IN):
        raise SystemExit(
            f'Falta {XLSX_IN}. Bajá "Volve production data.xlsx" (ver docstring) '
            'y guardalo ahí.')
    wb = openpyxl.load_workbook(XLSX_IN, read_only=True)
    ws = wb['Daily Production Data']
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    out: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        d = dict(zip(hdr, r))
        w = d['NPD_WELL_BORE_NAME']
        if w in WELLS and d['WELL_TYPE'] == 'OP':
            out[w].append(d)
    for w in out:
        out[w].sort(key=lambda d: d['DATEPRD'])
    return out


def monthly(days: list[dict]) -> dict:
    """Monthly medians of the daily diagnostics."""
    acc: dict[str, dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    for d in days:
        ym = d['DATEPRD'].strftime('%Y-%m')
        hrs = float(d['ON_STREAM_HRS'] or 0)
        oil = float(d['BORE_OIL_VOL'] or 0)
        wat = float(d['BORE_WAT_VOL'] or 0)
        bhp = float(d['AVG_DOWNHOLE_PRESSURE'] or 0)
        m = acc[ym]
        if hrs == 0 and bhp > P_SENSOR_MIN:
            m['p_res'].append(bhp)          # shut-in: reservoir-pressure proxy
        if hrs > 0 and oil > 0:
            m['q_oil'].append(oil * 24 / hrs)
            m['q_liq'].append((oil + wat) * 24 / hrs)
            if oil + wat > 0:
                m['wct'].append(wat / (oil + wat))
            if bhp > P_SENSOR_MIN:
                m['p_wf'].append(bhp)
    out = {'ym': sorted(acc)}
    for key, need in (('q_oil', 1), ('q_liq', 1), ('wct', 1), ('p_wf', 1),
                      ('p_res', MIN_SHUT_DAYS)):
        out[key] = [float(np.median(acc[ym][key])) if len(acc[ym][key]) >= need
                    else np.nan for ym in out['ym']]
    return out


def interp_nan(x: np.ndarray) -> np.ndarray:
    """Linear interpolation across NaN gaps (edges hold the nearest value)."""
    x = x.copy()
    ok = ~np.isnan(x)
    if ok.sum() < 2:
        return x
    idx = np.arange(len(x))
    x[~ok] = np.interp(idx[~ok], idx[ok], x[ok])
    return x


def analyze(m: dict) -> dict:
    n = len(m['ym'])
    t = np.arange(n) * 30.44  # days, monthly grid
    q_oil = np.array(m['q_oil'])
    p_res = interp_nan(np.array(m['p_res']))
    p_wf = np.array(m['p_wf'])
    q_liq = np.array(m['q_liq'])

    dd = p_res - p_wf
    with np.errstate(invalid='ignore', divide='ignore'):
        pi = np.where(dd > 5, q_liq / dd, np.nan)

    # Arps on the monthly oil rate, from its peak
    ok = ~np.isnan(q_oil)
    i0 = int(np.nanargmax(np.where(ok, q_oil, -1)))
    fit_mask = ok & (np.arange(n) >= i0)
    fit = None
    if fit_mask.sum() >= 12:
        tf = t[fit_mask] - t[i0]
        lnq = np.log(q_oil[fit_mask])
        try:
            f = lambda tt, qi, di, b: np.log(q_arps(tt, qi, di, b))
            popt, _ = curve_fit(f, tf, lnq, p0=[q_oil[i0], 0.001, 0.5],
                                bounds=([1e-3, 1e-6, 0.0], [q_oil[i0] * 5, 0.1, 1.2]),
                                maxfev=20000)
            resid = lnq - f(tf, *popt)
            r2 = 1 - np.sum(resid ** 2) / np.sum((lnq - lnq.mean()) ** 2)
            fit = {'qi': popt[0], 'di_anual': popt[1] * 365.25, 'b': popt[2],
                   'r2': float(r2), 'i0': i0}
        except RuntimeError:
            pass

    def trend(x, label_up, label_down, tol=0.15):
        """First-quarter vs last-quarter medians of the valid values."""
        v = x[~np.isnan(x)]
        if len(v) < 8:
            return 'sin datos suficientes'
        a, b_ = np.median(v[:len(v) // 4]), np.median(v[-len(v) // 4:])
        if b_ > a * (1 + tol):
            return f'{label_up} ({a:.0f} → {b_:.0f})'
        if b_ < a * (1 - tol):
            return f'{label_down} ({a:.0f} → {b_:.0f})'
        return f'estable ({a:.0f} → {b_:.0f})'

    wct = np.array(m['wct']) * 100
    diagnosis = {
        'presión de reservorio (proxy de cierres)': trend(p_res, 'SUBE', 'CAE'),
        'presión fluyente': trend(p_wf, 'sube', 'cae'),
        'drawdown pR - pwf': trend(dd, 'sube', 'cae'),
        'índice de productividad (líquido)': trend(pi, 'sube', 'cae'),
        'corte de agua (%)': trend(wct, 'SUBE', 'cae'),
        'caudal de líquido': trend(q_liq, 'sube', 'cae'),
        'caudal de petróleo': trend(q_oil, 'sube', 'CAE'),
    }
    return {'m': m, 't': t, 'p_res': p_res, 'pi': pi, 'fit': fit,
            'diagnosis': diagnosis}


def write_xlsx(res: dict[str, dict]) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Diagnóstico'
    for w, r in res.items():
        ws.append([w])
        for k, v in r['diagnosis'].items():
            ws.append(['', k, v])
        if r['fit']:
            f = r['fit']
            ws.append(['', 'Arps sobre el petróleo',
                       f"qi={f['qi']:.0f} sm3/d, Di={f['di_anual']:.2f}/año, "
                       f"b={f['b']:.2f}, R²={f['r2']:.2f}"])
        ws.append([])
    for w, r in res.items():
        sh = wb.create_sheet(w.replace('/', '_'))
        sh.append(['mes', 'q_oil_sm3d', 'q_liq_sm3d', 'wct', 'p_wf_bar',
                   'p_res_proxy_bar', 'PI_sm3d_bar'])
        m = r['m']
        for i, ym in enumerate(m['ym']):
            def v(x):
                return round(float(x[i]), 2) if not np.isnan(x[i]) else ''
            sh.append([ym, v(np.array(m['q_oil'])), v(np.array(m['q_liq'])),
                       v(np.array(m['wct'])), v(np.array(m['p_wf'])),
                       v(r['p_res']), v(r['pi'])])
    wb.save(OUT_XLSX)


def write_png(res: dict[str, dict]) -> None:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt

    fig, axes = plt.subplots(len(res), 4, figsize=(20, 4.5 * len(res)))
    for row, (w, r) in zip(np.atleast_2d(axes), res.items()):
        m, t = r['m'], r['t'] / 365.25
        q_oil, q_liq = np.array(m['q_oil']), np.array(m['q_liq'])
        wct = np.array(m['wct']) * 100

        ax = row[0]
        ax.plot(t, q_oil, '-', color='#2e7d32', label='petróleo')
        ax.plot(t, q_liq, '-', color='#666', lw=0.9, label='líquido')
        ax2 = ax.twinx()
        ax2.plot(t, wct, '-', color='#1565c0', lw=0.9)
        ax2.set_ylabel('corte de agua %', color='#1565c0')
        ax2.set_ylim(0, 100)
        ax.set_title(f'{w} · caudales y agua')
        ax.legend(fontsize=7)

        ax = row[1]
        ax.plot(t, r['p_res'], '-', color='#b71c1c', label='pR (proxy cierres)')
        ax.plot(t, np.array(m['p_wf']), '-', color='#ef9a9a', label='pwf fluyente')
        ax.set_title('presiones de fondo [bar]')
        ax.legend(fontsize=7)

        ax = row[2]
        ax.plot(t, r['pi'], '.', ms=4, color='#4527a0')
        ax.set_title('PI líquido [sm³/d/bar]')

        ax = row[3]
        ax.semilogy(t, q_oil, '.', ms=4, color='#888')
        if r['fit']:
            f = r['fit']
            tt = np.linspace(r['t'][f['i0']], r['t'][-1], 200)
            ax.semilogy(tt / 365.25, q_arps(tt - r['t'][f['i0']], f['qi'],
                                            f['di_anual'] / 365.25, f['b']),
                        color='#d62728',
                        label=f"Arps b={f['b']:.2f} R²={f['r2']:.2f}")
            ax.legend(fontsize=7)
        ax.set_title('petróleo, semilog + Arps')
    fig.suptitle('Volve, referencia del instructor — la presión no cae; el agua sube (años desde el arranque de cada pozo)')
    fig.tight_layout()
    fig.savefig(OUT_PNG, dpi=110)


def main() -> None:
    data = load()
    res = {}
    for w in WELLS:
        r = analyze(monthly(data[w]))
        res[w] = r
        print(f'\n=== {w}')
        for k, v in r['diagnosis'].items():
            print(f'  {k}: {v}')
        if r['fit']:
            f = r['fit']
            print(f"  Arps petróleo: qi={f['qi']:.0f}, Di={f['di_anual']:.2f}/año, "
                  f"b={f['b']:.2f}, R²={f['r2']:.2f}")
    write_xlsx(res)
    write_png(res)
    print(f'\n{OUT_XLSX}\n{OUT_PNG}')


if __name__ == '__main__':
    main()
