#!/usr/bin/env python
"""Reference consolidation of the six ARCH daily reports used in session 3.

The session-3 workshop asks a chatbot to consolidate six one-page PDFs of the
Agencia de Regulación y Control de Hidrocarburos (ARCH, Ecuador) into one Excel
workbook. This script does the same from the text layer (pdftotext -layout,
poppler), so the prompt on the page and this file are two copies of one
contract: change both or neither.

What it reads, per report:
- table 1: the 14 companies (anterior, día, incremento, estimado, cumplimiento),
  the private subtotal and the national total;
- table 2: only its TOTAL row (EP Petroecuador by block, anterior and día);
- table 3: wells producing, in workover or completion, and drilling;
- table 4: EP Petroecuador gas (its number format changes from day to day);
- the "Novedades principales" text, one row per company.

Checks it writes: companies against the report's totals, and each day's
"producción anterior" against the previous report's "producción del día" (the
volumes are preliminary and get revised; two operation days, 11 and 12
September, have no report of their own).

    uv run --with openpyxl python scripts/arch_consolidar.py

Outputs: public/descargas/arch_consolidado_referencia.xlsx (committed, plan B)
and scripts/_cache/arch_control.json (the answer key).
"""

from __future__ import annotations

import json
import os
import re
import subprocess

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
DESC = os.path.join(HERE, '..', 'public', 'descargas')
OUT = os.path.join(DESC, 'arch_consolidado_referencia.xlsx')
KEY = os.path.join(HERE, '_cache', 'arch_control.json')
DIAS = ['08', '09', '10', '11', '14', '15']

# Table-1 row labels as printed, in order; the second element is the type.
COMPANIAS = [
    ('EP PETROECUADOR', 'pública'),
    ('ANDES PETROLEUM ECUADOR LTD', 'privada'),
    ('CONSORCIO BLOQUE ESPEJO', 'privada'),
    ('CONSORCIO BLOQUE PERICO', 'privada'),
    ('ENAP SIPETROL S.A. - ENAP SIPEC', 'privada'),
    ('GENTE OIL ECUADOR PTE.LTD', 'privada'),
    ('GRAN TIERRA ENERGY COLOMBIA LLC', 'privada'),
    ('ORION ENERGY OCANOPB S.A.', 'privada'),
    ('ORIONOIL ER S.A', 'privada'),
    ('PCR-ECUADOR S.A.', 'privada'),
    ('PETROBELL S.A.', 'privada'),
    ('PETROLEOS DEL PACIFICO S.A. PACIFPETROL', 'privada'),
    ('PETROORIENTAL S.A.', 'privada'),
    ('PLUSPETROL ECUADOR B.V', 'privada'),
]
NUM = r'-?[\d.]+,\d+|-?[\d.]+'


def num_es(tok: str) -> float:
    """Spanish format: point for thousands, comma for decimals (23.607 = 23607)."""
    return float(tok.replace('.', '').replace(',', '.'))


def num_gas(tok: str) -> float:
    """Gas switches format between reports: 24.181,23 one day, 24191.03 the next."""
    if ',' in tok:
        return num_es(tok)
    if re.fullmatch(r'\d+\.\d{2}', tok):
        return float(tok)
    return num_es(tok)


def texto(dia: str) -> str:
    pdf = os.path.join(DESC, f'arch-reporte-diario-2026-09-{dia}.pdf')
    return subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True, check=True).stdout


def fila(lineas, etiqueta, n):
    """First n numbers after `etiqueta` on the line that starts with it."""
    for l in lineas:
        i = l.find(etiqueta)
        if i >= 0 and (etiqueta != 'EP PETROECUADOR' or 'PÚBLICA' in l):
            nums = re.findall(NUM, l[i + len(etiqueta):])
            if len(nums) >= n:
                return nums[:n]
    raise ValueError(f'no encontré la fila {etiqueta!r}')


def leer(dia: str) -> dict:
    t = texto(dia)
    lineas = t.splitlines()
    fechas = re.findall(r'(\d{1,2})/9/2026', '\n'.join(lineas[:12]))
    publicado, operacion = (f'2026-09-{int(d):02d}' for d in fechas[:2])
    t1 = []
    for nombre, tipo in COMPANIAS:
        ant, hoy, inc, est, cum = fila(lineas, nombre, 5)
        t1.append({'compania': nombre, 'tipo': tipo, 'anterior_bppd': num_es(ant), 'dia_bppd': num_es(hoy),
                   'incremento_bppd': num_es(inc), 'estimado_bppd': num_es(est), 'cumplimiento_pct': num_es(cum)})
    sub = [num_es(x) for x in fila(lineas, 'Subtotal Cías. Privadas', 5)]
    tot = [num_es(x) for x in fila(lineas, 'Total Nacional', 5)]
    bloques = next([num_es(x) for x in re.findall(NUM, l[l.rfind('TOTAL') + 5:])][:2]
                   for l in lineas if re.search(r'\s{20,}TOTAL\s+[\d.]+,\d+\s+[\d.]+,\d+\s*$', l))
    pozos = {}
    for etiqueta, clave in (('EMPRESA PÚBLICA', 'publica'), ('EMPRESAS PRIVADAS', 'privadas'), ('TOTAL NACIONAL', 'total')):
        prod, reac, perf = fila(lineas, etiqueta, 3)
        pozos[clave] = {'produccion': num_es(prod), 'reacondicionamiento': num_es(reac), 'perforacion': num_es(perf)}
    gas = next(re.findall(r'\d[\d.,]*', l.split('EP PETROECUADOR')[1])[:2]
               for l in lineas if 'EP PETROECUADOR' in l and 'PÚBLICA' not in l and re.search(r'EP PETROECUADOR\s+\d', l))
    nov = t.split('NOVEDADES PRINCIPALES:')[1]
    nov = re.sub(r'\s+', ' ', nov).strip()
    novedades = []
    for m in re.finditer(r'\*C[ÍI]A\.?\s*([^:]+?)(?:[:\-]\s)(.*?)(?=\*C[ÍI]A|$)', nov):
        novedades.append({'compania': m.group(1).strip(' .-'), 'texto': m.group(2).strip()})
    return {
        'publicado': publicado, 'operacion': operacion, 'companias': t1,
        'subtotal_privadas': dict(zip(('anterior_bppd', 'dia_bppd', 'incremento_bppd', 'estimado_bppd', 'cumplimiento_pct'), sub)),
        'total_nacional': dict(zip(('anterior_bppd', 'dia_bppd', 'incremento_bppd', 'estimado_bppd', 'cumplimiento_pct'), tot)),
        'ep_bloques_total': {'anterior_bppd': bloques[0], 'dia_bppd': bloques[1]},
        'pozos': pozos, 'gas_ep_mpcpd': {'anterior': num_gas(gas[0]), 'dia': num_gas(gas[1]), 'como_viene': gas},
        'novedades': novedades,
    }


def _siguiente(fecha: str) -> str:
    import datetime as dt
    return (dt.date.fromisoformat(fecha) + dt.timedelta(days=1)).isoformat()


def _anterior(fecha: str) -> str:
    import datetime as dt
    return (dt.date.fromisoformat(fecha) - dt.timedelta(days=1)).isoformat()


def controles(partes):
    out = []
    prev = None
    for p in partes:
        c = p['companias']
        s_priv = sum(x['dia_bppd'] for x in c if x['tipo'] == 'privada')
        s_tot = sum(x['dia_bppd'] for x in c)
        out.append({'operacion': p['operacion'], 'control': 'privadas suman el subtotal',
                    'calculado': round(s_priv, 2), 'parte': p['subtotal_privadas']['dia_bppd'],
                    'diferencia': round(s_priv - p['subtotal_privadas']['dia_bppd'], 2)})
        out.append({'operacion': p['operacion'], 'control': 'compañías suman el total nacional',
                    'calculado': round(s_tot, 2), 'parte': p['total_nacional']['dia_bppd'],
                    'diferencia': round(s_tot - p['total_nacional']['dia_bppd'], 2)})
        out.append({'operacion': p['operacion'], 'control': 'bloques suman EP Petroecuador',
                    'calculado': p['ep_bloques_total']['dia_bppd'], 'parte': c[0]['dia_bppd'],
                    'diferencia': round(p['ep_bloques_total']['dia_bppd'] - c[0]['dia_bppd'], 2)})
        if prev and _siguiente(prev['operacion']) != p['operacion']:
            out.append({'operacion': p['operacion'], 'control': f'sin parte de los días entre {prev["operacion"]} y {p["operacion"]}: la "producción anterior" de este parte es la única cifra del {_anterior(p["operacion"])}',
                        'calculado': None, 'parte': c[0]['anterior_bppd'], 'diferencia': 0})
        elif prev:
            for a, b in zip(prev['companias'], c):
                d = round(b['anterior_bppd'] - a['dia_bppd'], 2)
                if abs(d) > 0.005:
                    out.append({'operacion': p['operacion'], 'control': f'anterior de {b["compania"]} contra el parte previo ({prev["operacion"]})',
                                'calculado': a['dia_bppd'], 'parte': b['anterior_bppd'], 'diferencia': d})
        prev = p
    return out


def libro(partes, ctrl):
    wb = Workbook()
    ley = wb.active
    ley.title = 'Léeme'
    for i, t in enumerate([
        'Partes diarios de la ARCH (Ecuador), 8 al 15 de septiembre de 2026: planilla de referencia del curso (sesión 3).',
        'Fuente: Reporte diario preliminar de producción y operaciones, controlhidrocarburos.gob.ec. Volúmenes sujetos a revisión.',
        'Cada parte informa el día de operación anterior. Los días de operación 11 y 12 de septiembre no tienen parte propio.',
        'Números pasados a punto decimal. En el original: punto de miles y coma decimal; el gas cambia de formato según el día.',
        'Hoja Control: sumas contra los totales del parte, y la "producción anterior" de cada día contra el parte previo.',
    ], 1):
        ley.cell(row=i, column=1, value=t)
    ley.column_dimensions['A'].width = 130

    h = wb.create_sheet('Producción por compañía')
    h.append(['publicado', 'operacion', 'compania', 'tipo', 'anterior_bppd', 'dia_bppd', 'incremento_bppd', 'estimado_bppd', 'cumplimiento_pct'])
    for p in partes:
        for c in p['companias']:
            h.append([p['publicado'], p['operacion'], c['compania'], c['tipo'], c['anterior_bppd'], c['dia_bppd'],
                      c['incremento_bppd'], c['estimado_bppd'], c['cumplimiento_pct']])
    n = h.max_row

    r = wb.create_sheet('Resumen', 1)
    r.append(['operacion', 'EP Petroecuador (bppd)', 'Privadas (bppd)', 'Total nacional (bppd)', 'Total según el parte', 'Diferencia'])
    for i, p in enumerate(partes, 2):
        op = p['operacion']
        r.append([op,
                  f'=SUMIFS(\'Producción por compañía\'!$F$2:$F${n},\'Producción por compañía\'!$B$2:$B${n},A{i},\'Producción por compañía\'!$D$2:$D${n},"pública")',
                  f'=SUMIFS(\'Producción por compañía\'!$F$2:$F${n},\'Producción por compañía\'!$B$2:$B${n},A{i},\'Producción por compañía\'!$D$2:$D${n},"privada")',
                  f'=B{i}+C{i}', p['total_nacional']['dia_bppd'], f'=D{i}-E{i}'])
        for col in 'BCDEF':
            r[f'{col}{i}'].number_format = '#,##0.00'
    ch = LineChart()
    ch.title = 'Producción nacional por día de operación (bppd)'
    ch.add_data(Reference(r, min_col=4, min_row=1, max_row=len(partes) + 1), titles_from_data=True)
    ch.set_categories(Reference(r, min_col=1, min_row=2, max_row=len(partes) + 1))
    ch.height, ch.width = 8, 16
    r.add_chart(ch, 'H2')
    for col, w in zip('ABCDEF', (12, 22, 16, 22, 20, 12)):
        r.column_dimensions[col].width = w

    po = wb.create_sheet('Pozos y gas')
    po.append(['operacion', 'grupo', 'en producción', 'en reacondicionamiento o completación', 'en perforación'])
    for p in partes:
        for k, v in p['pozos'].items():
            po.append([p['operacion'], k, v['produccion'], v['reacondicionamiento'], v['perforacion']])
    po.append([])
    po.append(['operacion', 'gas EP Petroecuador (MPCPD)', 'como viene en el PDF'])
    for p in partes:
        po.append([p['operacion'], p['gas_ep_mpcpd']['dia'], p['gas_ep_mpcpd']['como_viene'][1]])

    nv = wb.create_sheet('Novedades')
    nv.append(['operacion', 'compania', 'texto original'])
    for p in partes:
        for x in p['novedades']:
            nv.append([p['operacion'], x['compania'], x['texto']])
    nv.column_dimensions['C'].width = 140

    k = wb.create_sheet('Control')
    k.append(['operacion', 'control', 'calculado', 'según el parte', 'diferencia'])
    for x in ctrl:
        k.append([x['operacion'], x['control'], x['calculado'], x['parte'], x['diferencia']])
    k.column_dimensions['B'].width = 80
    for ws in wb.worksheets[1:]:
        for c in ws[1]:
            c.font = Font(bold=True)
    wb.save(OUT)


def main() -> None:
    partes = [leer(d) for d in DIAS]
    ctrl = controles(partes)
    libro(partes, ctrl)
    os.makedirs(os.path.dirname(KEY), exist_ok=True)
    json.dump({'partes': partes, 'controles': ctrl}, open(KEY, 'w'), ensure_ascii=False, indent=1)
    print(f'{OUT}: {os.path.getsize(OUT):,} bytes')
    for p in partes:
        print(p['publicado'], 'op', p['operacion'], 'EP', p['companias'][0]['dia_bppd'], 'priv', p['subtotal_privadas']['dia_bppd'],
              'total', p['total_nacional']['dia_bppd'], 'bloques', p['ep_bloques_total']['dia_bppd'],
              'pozos', p['pozos']['total']['produccion'], 'gas', p['gas_ep_mpcpd']['dia'], p['gas_ep_mpcpd']['como_viene'][1],
              'novedades', len(p['novedades']))
    for x in ctrl:
        if abs(x['diferencia']) > 0.005 or x['calculado'] is None:
            print('  ', x['operacion'], x['control'], x['calculado'], x['parte'], x['diferencia'])


if __name__ == '__main__':
    main()
