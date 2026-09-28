#!/usr/bin/env python
"""Answer key for session 6: the verified facts of the Volve bundle.

Every Volve number printed on the session-6 page, the deck, the agent's
CLAUDE.md and the instructor README comes from this script, run on the FULL
bundle that scripts/build_volve_bundle.py leaves in scripts/_cache/volve/.
It is the hoja de respuestas of the live demo: it is never copied to the
agent's folder.

    uv run --with openpyxl --with pyproj python scripts/volve_checks.py
    ... --datos ~/volve-agente-vivo/datos   # same checks on a prepared folder

pyproj is optional (only the ED50 -> WGS84 shift needs it).

Output: a readable summary on stdout and scripts/_cache/volve_checks.json.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
FULL = os.path.join(HERE, '_cache', 'volve', 'completo')
LITE = os.path.join(HERE, '..', 'public', 'descargas', 'volve')
OUT = os.path.join(HERE, '_cache', 'volve_checks.json')

XLSX = 'produccion/Volve production data.xlsx'
PICKS = 'topes/Well_picks_Volve_v1.dat'
GRID = 'mapas/Hugin_Fm_Top.csv'
TRAJ = 'trayectorias/F-12_ACTUAL'
LAS_F12_IN = 'perfiles/15_9-F-12/WLC_PETRO_COMPUTED_INPUT_1.LAS'
LAS_F12_OUT = 'perfiles/15_9-F-12/WLC_PETRO_COMPUTED_OUTPUT_1.LAS'
LAS_F11B = 'perfiles/15_9-F-11 B/WLC_PETRO_COMPUTED_OUTPUT_1.LAS'
LAS_F14 = 'perfiles/15_9-F-14/NO_15_9-F-14_KLOGH_NEW.las'
SODIR_WELLS = 'sodir/sodir_volve_pozos.csv'
SODIR_PROD = 'sodir/sodir_volve_produccion_campo_mensual.csv'
LAS_LITE = '15_9-F-12_perfiles_2680-3600m.las'

NULL = -999.25
GR_PAD = 100.668998


def sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def read_las(path: str) -> dict:
    """Minimal LAS 2.0 reader: header items, curve mnemonics and the data rows."""
    header, curves, rows, section = {}, [], [], None
    for line in open(path, encoding='latin-1'):
        s = line.strip()
        if not s or s.startswith('#'):
            continue
        if s.startswith('~'):
            section = s[1].upper()
            continue
        if section in ('V', 'W', 'P'):
            m = re.match(r'([^.]+)\.(\S*)\s+(.*?)\s*:', s)
            if m:
                header[m.group(1).strip()] = m.group(3).strip()
        elif section == 'C':
            curves.append(s.split('.')[0].strip())
        elif section == 'A':
            rows.append([float(x) for x in s.split()])
    return {'header': header, 'curves': curves, 'rows': rows}


def curve(las: dict, name: str) -> list[tuple[float, float]]:
    j = las['curves'].index(name)
    return [(r[0], r[j]) for r in las['rows'] if r[j] != NULL]


def parse_picks(path: str) -> list[dict]:
    lines = open(path, encoding='latin-1').read().splitlines()
    rows, spans, names = [], None, None
    for i, line in enumerate(lines):
        if line.strip().startswith('Well name'):
            spans = [(m.start(), m.end()) for m in re.finditer(r'-+', lines[i + 1])]
            names = [line[a:b + 1].strip() for a, b in spans]
            continue
        if spans and line.startswith('  NO '):
            vals = []
            for k, (a, _) in enumerate(spans):
                end = spans[k + 1][0] if k + 1 < len(spans) else len(line)
                vals.append(line[a:end].strip())
            rows.append(dict(zip(names, vals)))
    return rows


def fnum(s: str):
    return float(s) if s not in ('', None) else None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--datos', default=FULL, help='carpeta con el paquete completo')
    args = ap.parse_args()
    D = os.path.expanduser(args.datos)
    P = lambda rel: os.path.join(D, rel)  # noqa: E731
    R: dict = {'datos': D}

    # ------------------------------------------------------------ files
    files = []
    for dirpath, _, names in os.walk(D):
        for n in sorted(names):
            p = os.path.join(dirpath, n)
            files.append({'archivo': os.path.relpath(p, D), 'bytes': os.path.getsize(p),
                          'sha256_12': sha256(p)[:12]})
    files.sort(key=lambda f: f['archivo'])
    R['archivos'] = files
    R['archivos_total_mb'] = round(sum(f['bytes'] for f in files) / 1e6, 1)

    # ------------------------------------------------------------ production
    import openpyxl
    wb = openpyxl.load_workbook(P(XLSX), read_only=True)
    mrows = list(wb['Monthly Production Data'].iter_rows(values_only=True))
    mdata = [r for r in mrows[2:] if r[0]]
    isnum = lambda v: isinstance(v, (int, float))  # noqa: E731
    per_well = defaultdict(lambda: defaultdict(float))
    months_oil = defaultdict(list)
    by_month = defaultdict(float)
    by_year = defaultdict(float)
    null_cells = Counter()
    hours_over = []
    for r in mdata:
        well, y, m = r[0], r[2], r[3]
        for key, v in zip(('horas', 'petroleo', 'gas', 'agua', 'gas_iny', 'agua_iny'), r[4:10]):
            if isnum(v):
                per_well[well][key] += v
            elif v in ('NULL', None):
                null_cells[key] += 1
        if isnum(r[5]) and r[5] > 0:
            months_oil[well].append(f'{y}-{m:02d}')
            by_month[(y, m)] += r[5]
            by_year[y] += r[5]
        days = (date(y + (m == 12), m % 12 + 1, 1) - date(y, m, 1)).days
        if isnum(r[4]) and r[4] > days * 24 + 1e-6:
            hours_over.append({'pozo': well, 'periodo': f'{y}-{m:02d}', 'horas': round(r[4], 2),
                               'horas_del_mes': days * 24})

    with open(P(SODIR_PROD), encoding='utf-8') as f:
        sprod = list(csv.DictReader(f))
    s_month = {(int(r['prfYear']), int(r['prfMonth'])): float(r['prfPrdOilNetMillSm3']) * 1e6 for r in sprod}
    s_year = defaultdict(float)
    for (y, _), v in s_month.items():
        s_year[y] += v
    s_oil = sum(s_month.values())
    s_water = sum(float(r['prfPrdProducedWaterInFieldMillSm3']) for r in sprod) * 1e6
    s_gas = sum(float(r['prfPrdGasNetBillSm3']) for r in sprod) * 1e9
    w_oil = sum(w['petroleo'] for w in per_well.values())
    w_water = sum(w['agua'] for w in per_well.values())
    w_gas = sum(w['gas'] for w in per_well.values())

    ranking = sorted(per_well, key=lambda w: -per_well[w]['petroleo'])
    R['produccion'] = {
        'filas_hoja_mensual': len(mdata),
        'pozos': len(per_well),
        'periodo_hoja_mensual': [min(f'{r[2]}-{r[3]:02d}' for r in mdata),
                                 max(f'{r[2]}-{r[3]:02d}' for r in mdata)],
        'celdas_NULL': dict(null_cells),
        'por_pozo': [{
            'pozo': w,
            'petroleo_sm3': round(per_well[w]['petroleo']),
            'participacion_petroleo_pct': round(100 * per_well[w]['petroleo'] / w_oil, 1),
            'agua_sm3': round(per_well[w]['agua']),
            'gas_sm3': round(per_well[w]['gas']),
            'agua_inyectada_sm3': round(per_well[w]['agua_iny']),
            'meses_con_petroleo': len(months_oil[w]),
            'primer_mes_petroleo': min(months_oil[w]) if months_oil[w] else None,
            'ultimo_mes_petroleo': max(months_oil[w]) if months_oil[w] else None,
        } for w in ranking],
        'horas_mensuales_mayores_al_mes': hours_over,
    }
    gaps = sorted(((k, by_month.get(k, 0.0) - s_month.get(k, 0.0)) for k in set(by_month) | set(s_month)),
                  key=lambda kv: kv[1])
    R['conciliacion'] = {
        'petroleo_pozos_sm3': round(w_oil),
        'petroleo_sodir_sm3': round(s_oil),
        'petroleo_pozos_millones_sm3': round(w_oil / 1e6, 2),
        'petroleo_sodir_millones_sm3': round(s_oil / 1e6, 2),
        'diferencia_sm3': round(w_oil - s_oil),
        'diferencia_pct': round(100 * (w_oil - s_oil) / s_oil, 2),
        'sodir_periodo': [f'{min(s_month)[0]}-{min(s_month)[1]:02d}', f'{max(s_month)[0]}-{max(s_month)[1]:02d}'],
        'primer_ultimo_mes_petroleo_pozos': [f'{min(by_month)[0]}-{min(by_month)[1]:02d}',
                                              f'{max(by_month)[0]}-{max(by_month)[1]:02d}'],
        'primer_ultimo_mes_petroleo_sodir': [
            '{}-{:02d}'.format(*min(k for k, v in s_month.items() if v > 0)),
            '{}-{:02d}'.format(*max(k for k, v in s_month.items() if v > 0))],
        'por_anio': [{'anio': y, 'pozos_sm3': round(by_year[y]), 'sodir_sm3': round(s_year[y]),
                      'diferencia_pct': round(100 * (by_year[y] - s_year[y]) / s_year[y], 1) if s_year[y] else None}
                     for y in sorted(set(by_year) | set(s_year)) if s_year[y] or by_year[y]],
        'meses_con_diferencia_mayor': [{'periodo': f'{k[0]}-{k[1]:02d}', 'diferencia_sm3': round(v)}
                                        for k, v in gaps[:5]],
        'meses_pozos_mayor_que_sodir': sum(1 for k in by_month if by_month[k] > s_month.get(k, 0) + 0.5),
        'agua_pozos_millones_sm3': round(w_water / 1e6, 2),
        'agua_sodir_millones_sm3': round(s_water / 1e6, 2),
        'agua_diferencia_pct': round(100 * (w_water - s_water) / s_water, 2),
        'gas_pozos_miles_millones_sm3': round(w_gas / 1e9, 3),
        'gas_neto_sodir_miles_millones_sm3': round(s_gas / 1e9, 3),
        'nota_gas': 'Sodir publica gas neto (vendible); el archivo por pozo, el gas producido. '
                    'No se concilian por definición.',
    }

    # ------------------------------------------------------------ daily sheet
    drows = list(wb['Daily Production Data'].iter_rows(values_only=True))
    dh = drows[0]
    dd = [dict(zip(dh, r)) for r in drows[1:]]
    over24 = [x for x in dd if (x['ON_STREAM_HRS'] or 0) > 24]
    neg = [x for x in dd if isnum(x['BORE_WAT_VOL']) and x['BORE_WAT_VOL'] < 0]
    d_oil = sum(x['BORE_OIL_VOL'] or 0 for x in dd)
    f5_op = [x for x in dd if x['NPD_WELL_BORE_NAME'] == '15/9-F-5' and x['WELL_TYPE'] == 'OP']
    R['diario'] = {
        'filas': len(dd),
        'periodo': [min(x['DATEPRD'] for x in dd).date().isoformat(),
                    max(x['DATEPRD'] for x in dd).date().isoformat()],
        'petroleo_diario_igual_al_mensual': abs(d_oil - w_oil) < 1,
        'filas_mas_de_24_horas': len(over24),
        'filas_con_25_horas': sum(1 for x in over24 if abs(x['ON_STREAM_HRS'] - 25) < 1e-9),
        'fechas_mas_de_24_horas': sorted({x['DATEPRD'].date().isoformat() for x in over24}),
        'nota_24h': 'Todas son el último domingo de octubre: el fin del horario de verano, un día de 25 horas.',
        'agua_negativa': [{'fecha': x['DATEPRD'].date().isoformat(), 'pozo': x['NPD_WELL_BORE_NAME'],
                           'agua_sm3': x['BORE_WAT_VOL']} for x in neg],
        'f5_dias_como_productor': len(f5_op),
        'f5_periodo_como_productor': [min(x['DATEPRD'] for x in f5_op).date().isoformat(),
                                      max(x['DATEPRD'] for x in f5_op).date().isoformat()] if f5_op else None,
        'f5_petroleo_sm3': round(per_well['15/9-F-5']['petroleo']),
    }
    for x in over24:
        assert x['DATEPRD'].month == 10 and x['DATEPRD'].weekday() == 6 and x['DATEPRD'].day >= 25

    # ------------------------------------------------------------ Sodir wells
    with open(P(SODIR_WELLS), encoding='utf-8') as f:
        swells = {r['wlbWellboreName']: r for r in csv.DictReader(f)}
    R['sodir_pozos'] = {
        'cantidad': len(swells),
        'datum': sorted({r['wlbGeodeticDatum'] for r in swells.values()}),
        'zona_utm': sorted({r['wlbUtmZone'] for r in swells.values()}),
        'por_proposito': dict(Counter(r['wlbPurpose'] for r in swells.values())),
    }

    # ------------------------------------------------------------ picks and grid
    picks = parse_picks(P(PICKS))
    grid = []
    with open(P(GRID), newline='') as f:
        for r in csv.DictReader(f):
            grid.append((int(r['IL']), int(r['XL']), float(r['X']), float(r['Y']), float(r['Z'])))
    ils = sorted({g[0] for g in grid})
    xls = sorted({g[1] for g in grid})
    idx = {(g[0], g[1]): g for g in grid}
    g0 = next(g for g in grid if (g[0] + 1, g[1]) in idx)
    g1 = idx[(g0[0] + 1, g0[1])]
    g2 = next(g for g in grid if (g[0], g[1] + 1) in idx)
    g3 = idx[(g2[0], g2[1] + 1)]
    R['grilla'] = {
        'puntos': len(grid),
        'inline': [ils[0], ils[-1]], 'crossline': [xls[0], xls[-1]],
        'z_min_m': round(min(g[4] for g in grid), 1), 'z_max_m': round(max(g[4] for g in grid), 1),
        'espaciado_inline_m': round(math.hypot(g1[2] - g0[2], g1[3] - g0[3]), 1),
        'espaciado_crossline_m': round(math.hypot(g3[2] - g2[2], g3[3] - g2[3]), 1),
        'este': [round(min(g[2] for g in grid)), round(max(g[2] for g in grid))],
        'norte': [round(min(g[3] for g in grid)), round(max(g[3] for g in grid))],
        'nota': 'Z es profundidad bajo el nivel del mar, positiva hacia abajo; en los topes, TVDSS es negativa.',
    }
    lite_grid = os.path.join(LITE, 'grilla_tope_hugin_raleada.csv')
    if os.path.exists(lite_grid):
        R['grilla']['puntos_raleada'] = sum(1 for _ in open(lite_grid)) - 1

    def nearest_z(x, y):
        best = min(grid, key=lambda g: (g[2] - x) ** 2 + (g[3] - y) ** 2)
        return best[4], math.hypot(best[2] - x, best[3] - y)

    R['topes'] = {
        'filas': len(picks),
        'pozos': len({p['Well name'] for p in picks}),
        'calificadores': dict(Counter(p['Qlf'] or '(normal)' for p in picks)),
        'superficies': len({p['Surface name'] for p in picks}),
    }
    entries = []
    for w in sorted(per_well):
        pick_name = 'NO ' + w + (' B' if w == '15/9-F-11' else '')
        hugin = [p for p in picks if p['Well name'] == pick_name and p['Surface name'] == 'Hugin Fm. VOLVE Top']
        hugin.sort(key=lambda p: float(p['MD']))
        if not hugin:
            entries.append({'pozo_produccion': w, 'pozo_topes': pick_name, 'entrada_hugin': None})
            continue
        h = hugin[0]
        z, dist = nearest_z(float(h['Easting']), float(h['Northing']))
        entries.append({'pozo_produccion': w, 'pozo_topes': pick_name, 'cruces_hugin': len(hugin),
                        'md_m': fnum(h['MD']), 'tvdss_m': fnum(h['TVDSS']), 'calificador': h['Qlf'],
                        'este_m': fnum(h['Easting']), 'norte_m': fnum(h['Northing']),
                        'grilla_z_m': round(z, 1), 'distancia_nodo_m': round(dist, 1),
                        'diferencia_grilla_menos_tope_m': round(z + fnum(h['TVDSS']), 1)})
    R['entrada_al_reservorio'] = entries

    # ------------------------------------------------------------ LAS
    f12 = read_las(P(LAS_F12_IN))
    f12o = read_las(P(LAS_F12_OUT))
    f11b = read_las(P(LAS_F11B))
    f14 = read_las(P(LAS_F14))
    gr = curve(f12, 'GR')
    pad = [d for d, v in gr if abs(v - GR_PAD) < 1e-6]
    last_real = max(d for d, v in gr if abs(v - GR_PAD) >= 1e-6)
    pad_run = [d for d in pad if d > last_real]
    td = float(swells['15/9-F-12']['wlbTotalDepth'])
    last_valid = {c: round(max(d for d, _ in curve(f12, c)), 2) for c in f12['curves'][1:] if curve(f12, c)}
    R['gr_f12'] = {
        'valor_constante_api': GR_PAD,
        'desde_m': round(min(pad_run), 2), 'hasta_m': round(max(pad_run), 2),
        'muestras': len(pad_run),
        'ultimo_valor_real_m': round(last_real, 2),
        'ultimo_valor_real_api': round(dict(gr)[last_real], 2),
        'profundidad_final_sodir_m': td,
        'metros_bajo_la_profundidad_final': round(max(pad_run) - td, 1),
        'muestras_bajo_la_profundidad_final': sum(1 for d in pad_run if d > td),
        'ultimo_dato_valido_por_curva_m': last_valid,
        'muestras_con_ese_valor_fuera_del_tramo': len(pad) - len(pad_run),
    }
    lite_las = os.path.join(LITE, LAS_LITE)
    if os.path.exists(lite_las):
        ll = read_las(lite_las)
        lgr = curve(ll, 'GR')
        R['gr_f12']['recorte_liviano'] = {
            'desde_m': ll['rows'][0][0], 'hasta_m': ll['rows'][-1][0], 'muestras': len(ll['rows']),
            'muestras_gr_constante': sum(1 for _, v in lgr if abs(v - GR_PAD) < 1e-6),
            'curvas': ll['curves'],
        }

    def constant_runs(las, min_len=21):
        """Runs of more than 20 consecutive samples with exactly the same non-null value."""
        out = []
        for j, name in enumerate(las['curves'][1:], start=1):
            run_start, prev, n = None, None, 0
            rows = las['rows']
            for k, r in enumerate(rows + [[None] * len(las['curves'])]):
                v = r[j]
                if v is not None and v != NULL and v == prev:
                    n += 1
                    continue
                if prev is not None and prev != NULL and n >= min_len:
                    out.append({'curva': name, 'valor': prev, 'desde_m': round(rows[run_start][0], 2),
                                'hasta_m': round(rows[k - 1][0], 2), 'muestras': n})
                run_start, prev, n = k, v, 1
        return out

    R['tramos_largos_f12_medidas'] = [x for x in constant_runs(f12, 101) if not x['curva'].endswith('_FLAG')]
    if os.path.exists(os.path.join(LITE, LAS_LITE)):
        R['tramos_largos_recorte_liviano'] = [x for x in constant_runs(read_las(os.path.join(LITE, LAS_LITE)), 101)]
    R['tramos_constantes'] = {}
    for las, rel in ((f12, LAS_F12_IN), (f12o, LAS_F12_OUT), (f11b, LAS_F11B), (f14, LAS_F14)):
        runs = constant_runs(las)
        R['tramos_constantes'][rel] = {
            'sin_banderas': [x for x in runs if not x['curva'].endswith('_FLAG')],
            'banderas': {c: sum(1 for x in runs if x['curva'] == c)
                         for c in sorted({x['curva'] for x in runs if x['curva'].endswith('_FLAG')})},
        }

    def las_summary(las, rel):
        depths = [r[0] for r in las['rows']]
        return {'archivo': rel, 'WELL': las['header'].get('WELL'), 'UWI': las['header'].get('UWI'),
                'curvas': las['curves'][1:], 'desde_m': round(min(depths), 2), 'hasta_m': round(max(depths), 2),
                'muestras': len(depths)}
    R['perfiles'] = [las_summary(f12, LAS_F12_IN), las_summary(f12o, LAS_F12_OUT),
                     las_summary(f11b, LAS_F11B), las_summary(f14, LAS_F14)]

    # ------------------------------------------------------------ F-11 vs F-11 B
    f11_picks = defaultdict(list)
    for p in picks:
        if p['Well name'].startswith('NO 15/9-F-11'):
            f11_picks[p['Well name']].append(p['Surface name'])
    R['f11'] = {
        'produccion_cargada_a': '15/9-F-11',
        'codigo_diario': sorted({x['WELL_BORE_CODE'] for x in dd if x['NPD_WELL_BORE_NAME'] == '15/9-F-11'}),
        'petroleo_sm3': round(per_well['15/9-F-11']['petroleo']),
        'meses_con_petroleo': [min(months_oil['15/9-F-11']), max(months_oil['15/9-F-11'])],
        'sodir': {n: {'proposito': r['wlbPurpose'], 'contenido': r['wlbContent'], 'inicio': r['wlbEntryDate'],
                      'terminacion': r['wlbCompletionDate'], 'profundidad_final_m': fnum(r['wlbTotalDepth'])}
                  for n, r in swells.items() if n.startswith('15/9-F-11')},
        'topes_por_rama': {k: {'filas': len(v), 'tiene_hugin': 'Hugin Fm. VOLVE Top' in v}
                           for k, v in sorted(f11_picks.items())},
        'perfil_del_paquete': f11b['header'].get('WELL'),
    }

    # ------------------------------------------------------------ F-14
    f14_depths = [r[0] for r in f14['rows']]
    f14_row = next(e for e in R['produccion']['por_pozo'] if e['pozo'] == '15/9-F-14')
    R['f14'] = {
        'petroleo_sm3': f14_row['petroleo_sm3'],
        'participacion_pct': f14_row['participacion_petroleo_pct'],
        'puesto_por_petroleo': ranking.index('15/9-F-14') + 1,
        'agua_sm3': f14_row['agua_sm3'],
        'puesto_por_agua': sorted(per_well, key=lambda w: -per_well[w]['agua']).index('15/9-F-14') + 1,
        'perfil_en_el_paquete': {'curvas': f14['curves'][1:], 'desde_m': min(f14_depths),
                                 'hasta_m': max(f14_depths), 'metros': round(max(f14_depths) - min(f14_depths), 1)},
        'profundidad_final_sodir_m': fnum(swells['15/9-F-14']['wlbTotalDepth']),
        'nota': 'En el espejo de origen, los perfiles interpretados de F-14 vienen en DLIS (binario); '
                'en LAS solo está la permeabilidad KLOGH_NEW.',
    }

    # ------------------------------------------------------------ coordinates F-12
    traj = [l.split() for l in open(P(TRAJ)) if re.match(r'\s*\d', l)]
    t0 = [float(v) for v in traj[0]]
    las_xy = (float(f12['header']['XCOORD']), float(f12['header']['YCOORD']))
    sod = swells['15/9-F-12']
    sod_xy = (float(sod['wlbEwUtm']), float(sod['wlbNsUtm']))
    seabed = next(p for p in picks if p['Well name'] == 'NO 15/9-F-12' and p['Surface name'] == 'Seabed')
    pick_xy = (float(seabed['Easting']), float(seabed['Northing']))
    d = lambda a, b: round(math.hypot(a[0] - b[0], a[1] - b[1]), 2)  # noqa: E731
    # LAS lat/lon vs Sodir lat/lon (spherical, fine at this scale)
    lat1, lon1 = float(f12['header']['LAT']), float(f12['header']['LON'])
    lat2, lon2 = float(sod['wlbNsDecDeg']), float(sod['wlbEwDecDeg'])
    dll = 6371000 * math.hypot(math.radians(lat1 - lat2), math.radians(lon1 - lon2) * math.cos(math.radians(lat1)))
    R['coordenadas_f12'] = {
        'las_encabezado_xy': las_xy,
        'trayectoria_primera_estacion': {'md_m': t0[0], 'xy': (t0[6], t0[7])},
        'sodir_utm': sod_xy,
        'tope_seabed_xy': pick_xy,
        'las_vs_trayectoria_m': d(las_xy, (t0[6], t0[7])),
        'las_vs_sodir_m': d(las_xy, sod_xy),
        'trayectoria_vs_sodir_m': d((t0[6], t0[7]), sod_xy),
        'trayectoria_vs_tope_seabed_m': d((t0[6], t0[7]), pick_xy),
        'las_latlon_vs_sodir_latlon_m': round(dll, 1),
        'trayectoria_profundidad_final_m': float(traj[-1][0]),
        'mesa_rotaria_las_m': float(f12['header']['ELEV']),
        'mesa_rotaria_sodir_m': float(sod['wlbKellyBushElevation']),
        'mesa_rotaria_topes_m': round(float(seabed['TVD']) + float(seabed['TVDSS']), 2),
    }
    try:
        from pyproj import Geod, Transformer
        t = Transformer.from_crs('EPSG:4230', 'EPSG:4326', always_xy=True)
        lon_w, lat_w = t.transform(lon2, lat2)
        az, _, dist = Geod(ellps='WGS84').inv(lon2, lat2, lon_w, lat_w)
        R['datum'] = {'ed50_a_wgs84_m': round(dist, 1), 'azimut_grados': round(az % 360),
                      'transformacion': 'pyproj EPSG:4230 a EPSG:4326, la transformación por defecto',
                      'nota': 'Otras transformaciones ED50 a WGS84 publicadas dan entre 113 y 117 m en este punto.'}
    except ImportError:
        R['datum'] = {'nota': 'pyproj no está: correr con --with pyproj para el corrimiento'}

    # ------------------------------------------------------------ F-12 well card (lite bundle)
    f12_row = next(e for e in R['produccion']['por_pozo'] if e['pozo'] == '15/9-F-12')
    R['ficha_f12'] = {
        'sodir': {k: sod[k] for k in ('wlbStatus', 'wlbPurpose', 'wlbContent', 'wlbEntryDate',
                                      'wlbCompletionDate', 'wlbTotalDepth', 'wlbFinalVerticalDepth',
                                      'wlbKellyBushElevation', 'wlbWaterDepth', 'wlbDrillingOperator')},
        'produccion': f12_row,
        'topes': [{'superficie': p['Surface name'], 'md_m': fnum(p['MD']), 'tvdss_m': fnum(p['TVDSS']),
                   'calificador': p['Qlf']} for p in picks if p['Well name'] == 'NO 15/9-F-12'],
    }

    # ------------------------------------------------------------ name variants
    def variants(base: str) -> dict:
        tok = base.split('/')[-1]  # e.g. 9-F-12
        pat = re.compile(r'F-' + re.escape(tok.split('F-')[1]) + r'(?![0-9])')
        out = defaultdict(set)
        for r in mdata:
            if r[0] == base:
                out['produccion mensual (pozo)'].add(r[0])
        for x in dd:
            if x['NPD_WELL_BORE_NAME'] == base:
                out['produccion diaria (NPD_WELL_BORE_NAME)'].add(x['NPD_WELL_BORE_NAME'])
                out['produccion diaria (WELL_BORE_CODE)'].add(x['WELL_BORE_CODE'])
        for p in picks:
            if pat.search(p['Well name']):
                out['topes (Well name)'].add(p['Well name'])
        for las, rel in ((f12, LAS_F12_IN), (f12o, LAS_F12_OUT), (f11b, LAS_F11B), (f14, LAS_F14)):
            if pat.search(rel):
                out['perfiles (carpeta)'].add(rel.split('/')[1])
                if pat.search(os.path.basename(rel)):
                    out['perfiles (archivo)'].add(os.path.basename(rel))
                out['perfiles (WELL)'].add(las['header'].get('WELL', ''))
                if las['header'].get('UWI'):
                    out['perfiles (UWI)'].add(las['header']['UWI'])
        if pat.search(TRAJ):
            out['trayectoria (archivo)'].add(os.path.basename(TRAJ))
        for n in swells:
            if pat.search(n):
                out['sodir (wlbWellboreName)'].add(n)
        flat = sorted({v for vs in out.values() for v in vs})
        same = [v for v in flat if same_wellbore(v) == base]
        return {'por_fuente': {k: sorted(v) for k, v in out.items()},
                'formas_del_mismo_pozo': same, 'cantidad': len(same),
                'otras_ramas': [v for v in flat if v not in same]}

    def same_wellbore(s: str) -> str:
        """Undo the spellings that name the SAME wellbore; sidetracks (A, B, T2, pilot) survive."""
        s = re.sub(r'\.las$', '', s, flags=re.I)
        s = re.sub(r'_(KLOGH_NEW|ACTUAL)$', '', s)
        s = re.sub(r'^NO[ _]', '', s).replace('15_9', '15/9')
        s = re.sub(r' A?H$', '', s)
        return s if s.startswith('15/9-') else '15/9-' + s

    R['nombres'] = {w: variants(w) for w in sorted(per_well)}

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(R, f, ensure_ascii=False, indent=2, default=str)

    # ------------------------------------------------------------ summary
    c = R['conciliacion']
    print(f"Paquete: {len(files)} archivos, {R['archivos_total_mb']} MB en {D}")
    print('\nPRODUCCIÓN POR POZO (hoja mensual)')
    for e in R['produccion']['por_pozo']:
        print(f"  {e['pozo']:<12} {e['petroleo_sm3']:>11,} Sm3  {e['participacion_petroleo_pct']:>5.1f}%  "
              f"agua {e['agua_sm3']:>11,}  {e['primer_mes_petroleo']} a {e['ultimo_mes_petroleo']}")
    print(f"\nCONCILIACIÓN  pozos {c['petroleo_pozos_sm3']:,} Sm3 ({c['petroleo_pozos_millones_sm3']} M)  "
          f"Sodir {c['petroleo_sodir_sm3']:,} ({c['petroleo_sodir_millones_sm3']} M)  "
          f"diferencia {c['diferencia_sm3']:,} ({c['diferencia_pct']}%)")
    print(f"  meses con más en pozos que en Sodir: {c['meses_pozos_mayor_que_sodir']}")
    for y in c['por_anio']:
        print(f"  {y['anio']}  {y['pozos_sm3']:>10,}  {y['sodir_sm3']:>10,}  {y['diferencia_pct']}%")
    print(f"  mayores diferencias: {c['meses_con_diferencia_mayor'][:3]}")
    print(f"  agua {c['agua_pozos_millones_sm3']} vs {c['agua_sodir_millones_sm3']} M ({c['agua_diferencia_pct']}%); "
          f"gas {c['gas_pozos_miles_millones_sm3']} vs neto {c['gas_neto_sodir_miles_millones_sm3']} miles de M")
    di = R['diario']
    print(f"\nDIARIO  {di['filas']:,} filas, {di['periodo']}; >24 h: {di['filas_mas_de_24_horas']} filas "
          f"({di['filas_con_25_horas']} con 25 h) en {di['fechas_mas_de_24_horas']}")
    print(f"  agua negativa: {di['agua_negativa']}")
    print(f"  F-5 como productor: {di['f5_dias_como_productor']} días {di['f5_periodo_como_productor']}, "
          f"{di['f5_petroleo_sm3']:,} Sm3")
    print(f"  horas mensuales mayores que el mes: {len(R['produccion']['horas_mensuales_mayores_al_mes'])}")
    g = R['gr_f12']
    print(f"\nGR F-12  {g['valor_constante_api']} API de {g['desde_m']} a {g['hasta_m']} m ({g['muestras']:,} muestras); "
          f"último real {g['ultimo_valor_real_m']} m; TD Sodir {g['profundidad_final_sodir_m']} m; "
          f"{g['muestras_bajo_la_profundidad_final']:,} muestras y {g['metros_bajo_la_profundidad_final']} m debajo")
    print(f"  último dato válido por curva: {g['ultimo_dato_valido_por_curva_m']}")
    if 'recorte_liviano' in g:
        print(f"  recorte liviano: {g['recorte_liviano']}")
    print(f"  tramos de más de 100 muestras en las curvas medidas de F-12: {R['tramos_largos_f12_medidas']}")
    print(f"  ídem en el recorte liviano: {R.get('tramos_largos_recorte_liviano')}")
    print('  tramos constantes (>20 muestras, sin banderas):')
    for rel, v in R['tramos_constantes'].items():
        for x in v['sin_banderas']:
            print(f"    {rel.split('/')[1]:<12} {x['curva']:<10} {x['valor']:>12} {x['desde_m']:>9} a {x['hasta_m']:>9} m  {x['muestras']:>5} muestras")
        if v['banderas']:
            print(f"    {rel.split('/')[1]:<12} banderas con tramos constantes: {v['banderas']}")
    f11 = R['f11']
    print(f"\nF-11  producción cargada a {f11['produccion_cargada_a']} ({f11['codigo_diario']}), "
          f"{f11['petroleo_sm3']:,} Sm3, {f11['meses_con_petroleo']}")
    for n, s in f11['sodir'].items():
        print(f"  Sodir {n:<14} {s['proposito']:<12} {s['inicio']} a {s['terminacion']}")
    print(f"  topes: {f11['topes_por_rama']}")
    f14r = R['f14']
    print(f"\nF-14  {f14r['petroleo_sm3']:,} Sm3 ({f14r['participacion_pct']}%), puesto {f14r['puesto_por_petroleo']} "
          f"en petróleo y {f14r['puesto_por_agua']} en agua; perfil: {f14r['perfil_en_el_paquete']}")
    co = R['coordenadas_f12']
    print(f"\nCOORDENADAS F-12  LAS vs trayectoria {co['las_vs_trayectoria_m']} m; LAS vs Sodir {co['las_vs_sodir_m']} m; "
          f"trayectoria vs Sodir {co['trayectoria_vs_sodir_m']} m; lat/lon LAS vs Sodir {co['las_latlon_vs_sodir_latlon_m']} m; "
          f"mesa rotaria LAS {co['mesa_rotaria_las_m']} / Sodir {co['mesa_rotaria_sodir_m']} / topes {co['mesa_rotaria_topes_m']} m")
    print(f"DATUM  {R['sodir_pozos']['datum']} zona {R['sodir_pozos']['zona_utm']}; {R['datum']}")
    print('\nENTRADA A HUGIN')
    for e in R['entrada_al_reservorio']:
        if e.get('md_m'):
            print(f"  {e['pozo_produccion']:<12} ({e['pozo_topes']}) MD {e['md_m']:>8,.1f}  TVDSS {e['tvdss_m']:>9,.1f}  "
                  f"grilla {e['grilla_z_m']:>7,.1f}  dif {e['diferencia_grilla_menos_tope_m']:>5}  {e['calificador'] or ''}")
        else:
            print(f"  {e['pozo_produccion']:<12} sin tope de Hugin en {e['pozo_topes']}")
    fi = R['ficha_f12']
    print(f"\nFICHA F-12  Sodir {fi['sodir']}")
    for t in fi['topes']:
        print(f"  {t['superficie']:<28} MD {t['md_m']:>8}  TVDSS {t['tvdss_m']}  {t['calificador']}")
    print(f"\nGRILLA {R['grilla']}")
    print(f"TOPES {R['topes']}")
    print('\nNOMBRES')
    for w, v in R['nombres'].items():
        print(f"  {w:<12} {v['cantidad']} formas: {v['formas_del_mismo_pozo']}  otras ramas: {v['otras_ramas']}")
    print(f'\nJSON: {OUT}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
